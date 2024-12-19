import os
import soundfile as sf
import audioread
from PIL import Image
import cv2
from docx import Document
import openpyxl
from gtts import gTTS
import pytesseract  # Para OCR
from InquirerPy import inquirer
from rich.console import Console
from rich.progress import Progress
from rich.table import Table
from rich.traceback import install
import matplotlib.pyplot as plt
import numpy as np
import librosa
import librosa.display

# Instala o traceback do rich para exibir erros de forma bonita
install()

console = Console()

def obter_extensoes_validas(tipo):
    if tipo == "imagem":
        return ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff', 'mp3']  # Adicionando mp3 como opção
    elif tipo == "áudio":
        return ['mp3', 'wav', 'aac', 'flac']
    elif tipo == "vídeo":
        return ['mp4', 'avi', 'mkv', 'mov']
    elif tipo == "documento":
        return ['pdf', 'docx', 'txt']
    elif tipo == "planilha":
        return ['xlsx', 'csv']
    return []

def converter_arquivo(tipo):
    try:
        arquivo = inquirer.text(message="Digite o caminho do arquivo a ser convertido:").execute()
        
        extensao = inquirer.select(
            message="Escolha a extensão para a qual deseja converter:",
            choices=obter_extensoes_validas(tipo)
        ).execute()
        
        pasta_destino = inquirer.text(message="Digite o caminho da pasta para salvar o arquivo convertido:").execute()
        
        if not os.path.exists(pasta_destino):
            os.makedirs(pasta_destino)

        novo_arquivo = os.path.join(pasta_destino, os.path.basename(arquivo).rsplit('.', 1)[0] + f'_convert.{extensao.lower()}')
        
        if os.path.exists(arquivo):
            console.print(f"\n[bold cyan]Iniciando a conversão:[/bold cyan] '{arquivo}' de [bold]{tipo}[/bold] para [bold]{extensao}[/bold]...")
            
            if tipo == "imagem":
                if extensao.lower() == "mp3":
                    # Converter imagem para áudio
                    img = Image.open(arquivo)
                    text = pytesseract.image_to_string(img)  # Extraindo texto da imagem
                    tts = gTTS(text=text, lang='pt')
                    tts.save(novo_arquivo)
                else:
                    img = Image.open(arquivo)
                    img.save(novo_arquivo)
            
            elif tipo == "áudio":
                with audioread.audio_open(arquivo) as input_file:
                    data = input_file.read_data()
                    sf.write(novo_arquivo, data, input_file.samplerate)
            
            elif tipo == "vídeo":
                cap = cv2.VideoCapture(arquivo)
                fourcc = cv2.VideoWriter_fourcc(*'XVID')
                out = cv2.VideoWriter(novo_arquivo, fourcc, 20.0, (640, 480))

                while cap.isOpened():
                    ret, frame = cap.read()
                    if not ret:
                        break
                    out.write(frame)

                cap.release()
                out.release()
            
            elif tipo == "documento":
                doc = Document(arquivo)
                doc.save(novo_arquivo)
            
            elif tipo == "planilha":
                wb = openpyxl.load_workbook(arquivo)
                wb.save(novo_arquivo)
            
            console.print(f"[green]Conversão concluída! Arquivo salvo como:[/green] '{novo_arquivo}'")
        else:
            console.print("[red]Erro:[/red] Arquivo não encontrado. Verifique o caminho e tente novamente.")
    except Exception as e:
        console.print(f"[red]Ocorreu um erro ao converter o arquivo:[/red] {str(e)}")

def converter_multiplos_arquivos(tipo):
    try:
        pasta = inquirer.text(message="Digite o caminho da pasta com os arquivos:").execute()
        
        if not os.path.exists(pasta):
            console.print("[red]Erro:[/red] Pasta não encontrada. Verifique o caminho e tente novamente.")
            return

        extensao_origem = inquirer.select(
            message="Escolha a extensão dos arquivos que deseja converter:",
            choices=obter_extensoes_validas(tipo)
        ).execute()

        extensoes_validas = obter_extensoes_validas(tipo)
        arquivos_encontrados = [f for f in os.listdir(pasta) if f.endswith(extensao_origem)]

        if not arquivos_encontrados:
            console.print("[red]Erro:[/red] Nenhum arquivo válido encontrado na pasta.")
            return

        extensao_destino = inquirer.select(
            message="Escolha a extensão para a qual deseja converter:",
            choices=extensoes_validas
        ).execute()

        novo_pasta = inquirer.text(message="Digite o caminho da pasta para salvar os arquivos convertidos:").execute()
        
        if not os.path.exists(novo_pasta):
            os.makedirs(novo_pasta)

        console.print(f"\n[bold cyan]Iniciando a conversão de múltiplos arquivos de [bold]{tipo}[/bold] na pasta:[/bold cyan] {pasta}")
        console.print(f"[bold cyan]Total de arquivos encontrados:[/bold cyan] {len(arquivos_encontrados)}")

        # Exibir tabela de arquivos
        table = Table(title="Arquivos a serem convertidos")
        table.add_column("Nome do Arquivo", style="cyan")
        for arquivo in arquivos_encontrados:
            table.add_row(arquivo)
        console.print(table)

        with Progress() as progress:
            task = progress.add_task("[cyan]Convertendo arquivos...", total=len(arquivos_encontrados))
            
            for arquivo in arquivos_encontrados:
                caminho_arquivo = os.path.join(pasta, arquivo)
                novo_arquivo = os.path.join(novo_pasta, arquivo.rsplit('.', 1)[0] + f'_convert.{extensao_destino.lower()}')
                
                console.print(f"\n[bold yellow]Convertendo:[/bold yellow] '{caminho_arquivo}' de [bold]{tipo}[/bold] para [bold]{extensao_destino}[/bold]...")
                
                if tipo == "imagem":
                    if extensao_destino.lower() == "mp3":
                        # Converter imagem para áudio
                        img = Image.open(caminho_arquivo)
                        text = pytesseract.image_to_string(img)  # Extraindo texto da imagem
                        tts = gTTS(text=text, lang='pt')
                        tts.save(novo_arquivo)
                    else:
                        img = Image.open(caminho_arquivo)
                        img.save(novo_arquivo)
                
                elif tipo == "áudio":
                    with audioread.audio_open(caminho_arquivo) as input_file:
                        data = input_file.read_data()
                        sf.write(novo_arquivo, data, input_file.samplerate)
                
                elif tipo == "vídeo":
                    cap = cv2.VideoCapture(caminho_arquivo)
                    fourcc = cv2.VideoWriter_fourcc(*'XVID')
                    out = cv2.VideoWriter(novo_arquivo, fourcc, 20.0, (640, 480))

                    while cap.isOpened():
                        ret, frame = cap.read()
                        if not ret:
                            break
                        out.write(frame)

                    cap.release()
                    out.release()
                
                elif tipo == "documento":
                    doc = Document(caminho_arquivo)
                    doc.save(novo_arquivo)
                
                elif tipo == "planilha":
                    wb = openpyxl.load_workbook(caminho_arquivo)
                    wb.save(novo_arquivo)
                
                console.print(f"[green]Arquivo convertido e salvo como:[/green] '{novo_arquivo}'")
                progress.update(task, advance=1)
    except Exception as e:
        console.print(f"[red]Ocorreu um erro ao converter os arquivos:[/red] {str(e)}")

def visualizar_espectro(arquivo_audio):
    # Carregar o arquivo de áudio
    y, sr = librosa.load(arquivo_audio)
    
    # Calcular o espectro
    D = librosa.amplitude_to_db(np.abs(librosa.stft(y)), ref=np.max)

    # Plotar o espectro
    plt.figure(figsize=(10, 4))
    librosa.display.specshow(D, sr=sr, x_axis='time', y_axis='log', cmap='coolwarm')
    plt.colorbar(format='%+2.0f dB')
    plt.title('Espectro de Frequência')
    plt.tight_layout()
    plt.show()

def main():
    while True:
        tipo = inquirer.select(
            message="Escolha o tipo de arquivo:",
            choices=["Imagem", "Áudio", "Vídeo", "Documento", "Planilha", "Visualizar Espectro", "Sair"]
        ).execute()

        if tipo == "Sair":
            break

        if tipo == "Visualizar Espectro":
            arquivo_audio = inquirer.text(message="Digite o caminho do arquivo de áudio:").execute()
            visualizar_espectro(arquivo_audio)
            continue

        is_multiple = inquirer.select(
            message="Você deseja converter um arquivo ou múltiplos arquivos?",
            choices=["Único", "Múltiplos"]
        ).execute()

        if is_multiple == "Único":
            converter_arquivo(tipo.lower())
        elif is_multiple == "Múltiplos":
            converter_multiplos_arquivos(tipo.lower())

if __name__ == "__main__":
    main()
